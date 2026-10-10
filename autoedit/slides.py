"""발표 슬라이드(pptx)를 이미지로 바꾸고, 각 슬라이드를 그 내용을 말하는 순간에 맞춘다.

슬라이드 문구는 대본을 줄여 쓴 것이라 말과 글자가 똑같지 않다. 그래서 글자 두 개씩 묶은 조각(바이그램)이
얼마나 겹치는지로 비교하고, 슬라이드 순서와 말하는 순서가 같다는 점을 이용해 전체를 한 번에 맞춘다.
"""
import re
import shutil
import subprocess
import urllib.request
from dataclasses import dataclass
from pathlib import Path

from .cuts import Sentence

FONT_DIR = Path.home() / ".local" / "share" / "fonts" / "autoedit"
# 슬라이드에 자주 쓰는 무료 폰트. LibreOffice는 웹폰트(woff)를 못 읽어서 받은 뒤 ttf로 바꿔 설치한다.
NOONNU = "https://raw.githubusercontent.com/projectnoonnu"
KNOWN_FONTS = {
    "SB Aggro": [f"{NOONNU}/noonfonts_2108/main/SBAggro{w}.woff" for w in "BML"],
    "TmonMonsori": [f"{NOONNU}/noonfonts_two/master/TmonMonsori.woff"],
}
SCORE_MIN = 0.12  # 이보다 덜 겹치면 맞는 문장이 없다고 본다
SKIP_SLIDE = -0.05  # 슬라이드를 건너뛰는 비용 (조금이라도 맞으면 쓰는 쪽을 택한다)


@dataclass
class Slide:
    no: int        # 1부터
    path: Path     # 1920x1080 PNG
    text: str      # 본문 글자 (머리말/쪽번호 제외)


def _install_fonts(pptx: Path):
    """pptx에 쓰인 폰트 중 아는 무료 폰트가 시스템에 없으면 받아서 설치한다."""
    xml = subprocess.run(["unzip", "-p", str(pptx), "ppt/slides/*.xml"], capture_output=True).stdout.decode("utf-8", "ignore")
    used = set(re.findall(r'typeface="([^"]+)"', xml))
    have = subprocess.run(["fc-list", ":", "family"], capture_output=True, text=True).stdout
    todo = [f for f in KNOWN_FONTS if f in used and f not in have]
    if not todo:
        return
    from fontTools.ttLib import TTFont

    FONT_DIR.mkdir(parents=True, exist_ok=True)
    for fam in todo:
        for url in KNOWN_FONTS[fam]:
            try:
                with urllib.request.urlopen(url, timeout=30) as r:
                    data = r.read()
                woff = FONT_DIR / Path(url).name
                woff.write_bytes(data)
                f = TTFont(str(woff))
                f.flavor = None
                f.save(str(woff.with_suffix(".ttf")))
                woff.unlink()
            except Exception as e:
                print(f"  슬라이드 폰트 {fam}를 받지 못했습니다 ({e}). 다른 글꼴로 대신 그려집니다.")
    _patch_aggro_bold()
    subprocess.run(["fc-cache", "-f"], capture_output=True)


def _patch_aggro_bold():
    """배포되는 어그로체 B 웹폰트는 '볼' 글자가 꽉 찬 사각형으로 깨져 있어서 M의 글자로 바꿔 넣는다."""
    b, m = FONT_DIR / "SBAggroB.ttf", FONT_DIR / "SBAggroM.ttf"
    if not (b.exists() and m.exists()):
        return
    from fontTools.ttLib import TTFont

    fb, fm = TTFont(str(b)), TTFont(str(m))
    for ch in "볼":
        gb, gm = fb.getBestCmap().get(ord(ch)), fm.getBestCmap().get(ord(ch))
        if gb and gm:
            fb["glyf"][gb] = fm["glyf"][gm]
            fb["hmtx"][gb] = fm["hmtx"][gm]
    fb.save(str(b))


def _body_text(pptx: Path) -> list[str]:
    """슬라이드마다 가장 큰 글자들(본문)만 모은다. 머리말, 쪽번호, 논문 원문 인용 같은 작은 글씨는 뺀다."""
    from pptx import Presentation

    out = []
    for s in Presentation(str(pptx)).slides:
        runs = []
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    size = r.font.size.pt if r.font.size else 18
                    runs.append((size, r.text))
        big = max((sz for sz, _ in runs), default=0)
        out.append(" ".join(t for sz, t in runs if sz >= min(24, big * 0.5)))
    return out


def load(pptx: Path, work: Path) -> list[Slide]:
    """pptx → PDF → 1920x1080 PNG. 이미 만들어 둔 이미지가 있으면 다시 쓴다."""
    folder = work / "slides"
    pngs = sorted(folder.glob("s-*.png"))
    src = folder / "src"
    if not pngs or not src.exists() or src.read_text(encoding="utf-8") != str(pptx):
        if not shutil.which("soffice"):
            raise RuntimeError("슬라이드를 그리려면 LibreOffice(soffice)가 필요합니다.")
        _install_fonts(pptx)
        shutil.rmtree(folder, ignore_errors=True)
        folder.mkdir(parents=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(folder), str(pptx)],
                       check=True, capture_output=True, timeout=600)
        pdf = folder / (pptx.stem + ".pdf")
        subprocess.run(["pdftoppm", "-png", "-scale-to-x", "1920", "-scale-to-y", "1080", str(pdf), str(folder / "s")],
                       check=True)
        src.write_text(str(pptx), encoding="utf-8")
        pngs = sorted(folder.glob("s-*.png"))
    texts = _body_text(pptx)
    return [Slide(i + 1, p, texts[i] if i < len(texts) else "") for i, p in enumerate(pngs)]


def _grams(text: str) -> set[str]:
    t = re.sub(r"[^가-힣A-Za-z0-9]", "", text)
    return {t[i:i + 2] for i in range(len(t) - 1)}


def _overlap(slide: set[str], speech: set[str]) -> float:
    return len(slide & speech) / len(slide) if slide else 0.0


def match(slides: list[Slide], sentences: list[Sentence], words: list[dict],
          chapters: list[dict] = ()) -> list[dict]:
    """슬라이드마다 원본 영상의 시각(초)을 정한다.

    1. 순서를 지키면서 글자 겹침 합이 가장 큰 (슬라이드, 문장) 짝을 찾는다. 한 문장에 여러 장이 맞을 수 있다.
    2. 챕터 제목과 같은 슬라이드(PART 1 같은 표지)는 그 챕터가 시작하는 순간에 고정한다.
    3. 말과 글자가 거의 안 겹치는 슬라이드(숫자, 논문 인용)는 앞뒤 슬라이드 사이에 고르게 끼운다.
    반환: [{"slide": Slide, "t": 원본 초, "matched": "text" | "chapter" | "between"}] (시간 순)
    """
    sg = [_grams(s.text) for s in slides]
    # 앞뒤 문장을 조금 붙여서 비교한다 (슬라이드 한 장이 문장 경계를 넘는 경우)
    ctx = [_grams(" ".join(x.text for x in sentences[max(0, j - 1):j + 2])) for j in range(len(sentences))]
    own = [_grams(x.text) for x in sentences]
    score = [[0.6 * _overlap(sg[i], own[j]) + 0.4 * _overlap(sg[i], ctx[j]) for j in range(len(sentences))]
             for i in range(len(slides))]

    n, m = len(slides), len(sentences)
    # best[i][j]: 슬라이드 i장까지 정했고 마지막으로 쓴 문장이 j번째 이하일 때의 최대 점수
    best = [[0.0] * (m + 1) for _ in range(n + 1)]
    back: list[list] = [[None] * (m + 1) for _ in range(n + 1)]
    for i in range(1, n + 1):
        for j in range(0, m + 1):
            cands = [(best[i - 1][j] + SKIP_SLIDE, ("skip", j))]
            if j > 0:
                cands.append((best[i][j - 1], ("left", j - 1)))
                sc = score[i - 1][j - 1]
                if sc >= SCORE_MIN:
                    cands.append((best[i - 1][j] + sc, ("take", j)))  # 같은 문장에 여러 장 허용
            best[i][j], back[i][j] = max(cands, key=lambda x: x[0])
    times: dict[int, tuple[float, str]] = {}
    took: dict[int, float] = {}
    i, j = n, m
    while i > 0:
        kind, pj = back[i][j]
        if kind == "left":
            j = pj
            continue
        if kind == "take":
            times[i - 1] = (_start_time(sg[i - 1], sentences, j - 1, words), "text")
            took[i - 1] = score[i - 1][j - 1]
        i -= 1
        j = pj

    # 챕터 표지 슬라이드: 챕터 이름과 많이 겹치는 슬라이드. 챕터 순서와 슬라이드 순서가 어긋나는 짝은 버린다
    by_idx = {s.idx: s for s in sentences}
    anchors = []  # (챕터 순서, 슬라이드 번호, 시작 시각)
    for ci, c in enumerate(chapters):
        lg = _grams(c.get("label", ""))
        if not lg or c["from"] not in by_idx:
            continue
        k = max(range(n), key=lambda k: _overlap(lg, sg[k]))
        t0 = by_idx[c["from"]].start
        # 말과 잘 맞은 슬라이드가 챕터 시작에서 먼 곳에 있으면 그대로 둔다 (챕터 이름과 비슷한 문구가 본문에 나오는 경우)
        near = k in times and abs(times[k][0] - t0) < 15
        if _overlap(lg, sg[k]) >= 0.25 and (took.get(k, 0) < 0.5 or near):
            anchors.append((ci, k, t0))
    for _, k, t in _increasing(anchors):
        times[k] = (t, "chapter")
    # 앞 슬라이드보다 이른 시각으로 잡힌 것은 잘못 맞은 것이라 버린다 (챕터 표지를 기준으로)
    fixed = sorted((k for k, v in times.items() if v[1] == "chapter"))
    for k in list(times):
        if times[k][1] == "chapter":
            continue
        before = [times[f][0] for f in fixed if f < k]
        after = [times[f][0] for f in fixed if f > k]
        if (before and times[k][0] < max(before)) or (after and times[k][0] > min(after)):
            del times[k]
    last = -1.0
    for k in sorted(times):
        if times[k][0] < last:
            del times[k]
        else:
            last = times[k][0]

    # 맞는 문장이 없는 슬라이드는 앞뒤 사이에 고르게
    end = sentences[-1].end if sentences else 0.0
    k = 0
    while k < n:
        if k in times:
            k += 1
            continue
        run_end = k
        while run_end < n and run_end not in times:
            run_end += 1
        t0 = times[k - 1][0] if k > 0 else (sentences[0].start if sentences else 0.0)
        t1 = times[run_end][0] if run_end < n else end
        cnt = run_end - k
        for q in range(cnt):
            times[k + q] = (t0 + (t1 - t0) * (q + 1) / (cnt + 1), "between")
        k = run_end
    return [{"slide": slides[k], "t": round(times[k][0], 2), "matched": times[k][1]} for k in sorted(times)]


def _increasing(anchors: list[tuple]) -> list[tuple]:
    """챕터 순서대로 놓았을 때 슬라이드 번호도 커지는 가장 긴 부분열."""
    best: list[list[tuple]] = []
    for a in anchors:
        prev = max((b for b in best if b[-1][1] < a[1]), key=len, default=[])
        best.append(prev + [a])
    return max(best, key=len, default=[])


def _start_time(slide: set[str], sentences: list[Sentence], j: int, words: list[dict]) -> float:
    """문장 j(과 다음 문장) 안에서 슬라이드 내용이 시작되는 단어의 시각.

    슬라이드 글자 수만큼의 단어 창을 밀면서 슬라이드 바이그램을 가장 많이 덮는 위치를 찾고,
    거의 같으면 앞쪽을 택한다.
    """
    ids = sentences[j].word_ids + (sentences[j + 1].word_ids if j + 1 < len(sentences) else [])
    if not ids or not slide:
        return sentences[j].start
    need = len(slide) + 1  # 대략의 글자 수
    scores = []
    for k in range(max(1, len(ids) - 2)):
        text, q = "", k
        while q < len(ids) and len(re.sub(r"\s", "", text)) < need * 1.3:
            text += words[ids[q]]["w"]
            q += 1
        scores.append(_overlap(slide, _grams(text)))
    top = max(scores)
    k = next(k for k, v in enumerate(scores) if v >= top * 0.9)
    return words[ids[k]]["s"]


def to_sentence(t: float, sentences: list[Sentence]) -> tuple[int, float]:
    """원본 시각을 (문장 번호, 문장 시작 후 초)로."""
    cur = sentences[0]
    for s in sentences:
        if s.start <= t + 1e-6:
            cur = s
    return cur.idx, round(max(0.0, t - cur.start), 2)
